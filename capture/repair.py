"""What Excel's repair does to a cell whose number format Excel refuses in a file.

Usage: ``uv run python -m capture.repair``. The one interactive capture: Excel refuses a file
holding such a format (``Parameter error (-50)`` with alerts off, CAPTURING.md section 2 item
9), so this opens two workbooks with alerts ON and a person accepts Excel's repair dialogs. The
first workbook holds every refused format known, one per row, between rows of valid custom
formats as controls; the second holds a single refused format, to see whether the repair treats
it the same alone. After each repair the script reads every case cell's displayed text and
number format. Writes ``tests/captures/repair.layout.tsv``; CAPTURING.md section 2 item 33.

No clipboard. The workbooks are written inside Excel's container and closed without saving.
"""

from __future__ import annotations

import csv
import datetime as dt
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE

from capture.build import WORKBOOK_DIR

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tests" / "captures" / "repair.layout.tsv"
SCRIPT = Path(__file__).with_name("repair.applescript")
US = "\x1f"
VALUES: tuple[object, ...] = (1234.5678, -0.5, 0, "abc")

REFUSED: tuple[str, ...] = (
    "0.00e+00", "0.00e-00", "##0.0e+0", "0.00E+00;0.00e+00", "0;0;0;0;0", "[Red", '0.00"',
    "[Red][Blue]0", "[>1][<5]0", "@;@", "0;@;0", "yyyy 0", "0 yyyy", "0.00E+", "# ?/",
    "[hh]:mm:[ss]", "General General",
    "[Color57]0", "[Color0]0", "[]0", "[abc]0", "ss.0000", "yyyy @", "0 @", "@ 0", "# ?/?/?",
    "[>1]0;[>2]0;[>3]0", "[>1]0;[>2]0;[>3]0;@", "0;0;0;[>1]@",
)  # fmt: skip
CONTROLS: tuple[str, ...] = ("0.000", "[Red]0.00", '"x"0', "#,##0.0", "m/d/yy")


def rows_all() -> list[tuple[str, str]]:
    """(tag, format) per row: a control, then each refused format followed by a control."""
    out = [("control", CONTROLS[0])]
    for k, fmt in enumerate(REFUSED):
        out.append(("refused", fmt))
        out.append(("control", CONTROLS[(k + 1) % len(CONTROLS)]))
    return out


def rows_single() -> list[tuple[str, str]]:
    return [("control", CONTROLS[0]), ("refused", "0.00e+00"), ("control", CONTROLS[1])]


def build(name: str, rows: list[tuple[str, str]]) -> Path:
    BUILTIN_FORMATS_REVERSE.clear()
    wb = Workbook()
    ws = wb.active
    ws.title = "repair"
    for r, (tag, fmt) in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=f"{tag} {fmt}").data_type = "s"
        for c, value in enumerate(VALUES, 2):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.number_format = fmt
    WORKBOOK_DIR.mkdir(parents=True, exist_ok=True)
    path = (WORKBOOK_DIR / name).resolve()
    wb.save(path)
    return path


def run(path: Path, rows: list[tuple[str, str]]) -> tuple[str, list[list[str]]]:
    print(f"opening {path.name} with alerts on: accept Excel's repair dialogs")
    proc = subprocess.run(
        ["osascript", str(SCRIPT), str(path), str(len(rows)), str(len(VALUES))],
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode != 0:
        sys.exit(f"osascript failed: {proc.stderr.strip()}")
    lines = proc.stdout.rstrip("\n").split("\n")
    workbook = lines[0].split(US)[1]
    return workbook, [line.split(US)[1:] for line in lines[1:]]


def main() -> None:
    """Build both workbooks, open each with alerts on, write what the repair left."""
    out: list[list[str]] = []
    for name, rows in (("repair_all.xlsx", rows_all()), ("repair_single.xlsx", rows_single())):
        repaired, cells = run(build(name, rows), rows)
        print(f"  repaired workbook: {repaired}")
        for (tag, fmt), fields in zip(rows, cells, strict=True):
            for k, value in enumerate(VALUES):
                out.append([name, tag, repr(fmt), repr(value), fields[2 * k], fields[2 * k + 1]])
    with OUTPUT.open("w", encoding="utf-8", newline="") as f:
        f.write(
            "# suite: repair (capture.repair); Excel for Mac, en_US; alerts on, repair accepted\n"
        )
        f.write(f"# captured: {dt.date.today().isoformat()}; workbook_font: Calibri 11\n")
        writer = csv.writer(f, dialect="excel-tab", lineterminator="\n")
        writer.writerow(["workbook", "tag", "format", "value_repr", "text", "format_after"])
        writer.writerows(out)
    print(f"wrote {OUTPUT}: {len(out)} rows")


if __name__ == "__main__":
    main()
