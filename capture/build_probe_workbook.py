"""Build the workbooks of the first file probes (CAPTURING.md section 2 items 8 and 9).

Usage: ``uv run python -m capture.build_probe_workbook``. Writes, into ``probes/`` under Excel's
container folder (``capture.build.WORKBOOK_DIR``, so Excel opens them without a permission
dialog):

- ``probe_halfdown.xlsx`` and ``probe_invalid.xlsx``: the half-down candidates and the
  dialog-rejected formats, written by openpyxl so that Excel meets them in a file rather than
  through the Format Cells dialog; read with ``capture/probe_file.applescript <path>``;
- ``invalid/inv_NN.xlsx``: one workbook per dialog-rejected format, so a file Excel refuses
  names its culprit; read with ``capture/probe_file_many.applescript <directory>``.

Layout: A = value with format applied, B = case id, C = repr(value), D = repr(format), in
sheet ``probe``. Later probes of refused formats use ``capture/invalid.py``; the rounding
probe became the ``rounding_15digit`` suite.
"""

from openpyxl import Workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE

from capture.build import WORKBOOK_DIR

BUILD = WORKBOOK_DIR / "probes"

HALF_DOWN_VALUES = [
    123451234512344.5,
    123451234512345.5,
    123451234512345.51,
    123451234512345.49,
    1234512345123455.5,
    12345123451234.5,
    # extra pairs to separate models: 16 sig digits ending in 6/4, and 17 digits
    123451234512345.6,
    123451234512345.4,
    1234512345123454.5,
    1234512345123456.5,
    99999999999999.5,
    999999999999999.5,
    2.5,
    0.125,
    1.0625,
    1234.5,
]
HALF_DOWN_FORMATS = ["0", "General", "0.0"]

INVALID_FORMATS = [
    "0.00e+00",  # lower-case exponent
    "0.00e-00",
    "##0.0e+0",
    "0.00E+00;0.00e+00",  # lower-case only in section 2
    "0;0;0;0;0",  # five sections
    "*-0*-",  # two fills
    "[Red",  # unclosed bracket
    '0.00"',  # unclosed quote
    "[Red][Blue]0",  # two colors
    "[>1][<5]0",  # two conditions
    "@;@",  # text in two sections
    "0;@;0",  # text placeholder in the middle
    "yyyy 0",  # date and digit codes mixed
    "0 yyyy",
    "0.00E+",  # exponent without digits
    "# ?/",  # fraction without denominator
    "e+",  # bare
    "[hh]:mm:[ss]",  # two elapsed
    "General General",
    "0.0.0",  # two decimal points
    "0%E+0",
]
INVALID_VALUES = [1234.5678, -0.5, 0, 45000.75, "abc"]


def write(name: str, prefix: str, formats: list[str], values: list[object]) -> None:
    BUILTIN_FORMATS_REVERSE.clear()
    wb = Workbook()
    ws = wb.active
    ws.title = "probe"
    row = 0
    for fmt in formats:
        for value in values:
            row += 1
            c = ws.cell(row=row, column=1, value=value)
            c.number_format = fmt
            ws.cell(row=row, column=2, value=f"{prefix}-{row}")
            ws.cell(row=row, column=3, value=repr(value))
            ws.cell(row=row, column=4, value=repr(fmt))
    ws.column_dimensions["A"].width = 60
    BUILD.mkdir(parents=True, exist_ok=True)
    out = BUILD / name
    wb.save(out)
    print(f"wrote {out} rows={row}")


def write_one_per_format() -> None:
    """One workbook per dialog-rejected format, so an unopenable file names its culprit."""
    for k, fmt in enumerate(INVALID_FORMATS, 1):
        BUILTIN_FORMATS_REVERSE.clear()
        wb = Workbook()
        ws = wb.active
        ws.title = "probe"
        for row, value in enumerate(INVALID_VALUES, 1):
            c = ws.cell(row=row, column=1, value=value)
            c.number_format = fmt
            ws.cell(row=row, column=2, value=f"inv{k}-{row}")
            ws.cell(row=row, column=3, value=repr(value))
            ws.cell(row=row, column=4, value=repr(fmt))
        ws.column_dimensions["A"].width = 60
        (BUILD / "invalid").mkdir(parents=True, exist_ok=True)
        wb.save(BUILD / "invalid" / f"inv_{k:02d}.xlsx")
    print(f"wrote {len(INVALID_FORMATS)} per-format workbooks")


def main() -> None:
    """Write every probe workbook."""
    write("probe_halfdown.xlsx", "hd", HALF_DOWN_FORMATS, HALF_DOWN_VALUES)
    write("probe_invalid.xlsx", "inv", INVALID_FORMATS, INVALID_VALUES)
    write_one_per_format()


if __name__ == "__main__":
    main()
