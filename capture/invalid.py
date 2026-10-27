"""Check which candidate format strings Excel refuses to open from a file.

Usage: ``uv run python -m capture.invalid``. Builds one workbook per candidate (five values each)
into Excel's container, runs ``capture/probe_file_many.applescript`` on the directory, and prints
one line per candidate: ``refused`` or the five displayed texts. Results are recorded by hand in
``CAPTURING.md`` section 2 item 9; refused strings become unit tests in ``tests/test_parser.py``
and rendered ones go into the ``odd_formats`` suite in ``capture/cases.py``.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE

from capture.build import WORKBOOK_DIR

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "capture" / "probe_file_many.applescript"
VALUES: list[object] = [1234.5678, -0.5, 0, 45000.75, "abc"]

CANDIDATES: dict[str, str] = {
    "color-index-57": "[Color57]0",
    "color-index-0": "[Color0]0",
    "color-with-space": "[Color 3]0",
    "empty-bracket": "[]0",
    "unknown-bracket": "[abc]0",
    "four-fraction-second-digits": "ss.0000",
    "at-in-date-section": "yyyy @",
    "at-after-digits": "0 @",
    "at-before-digits": "@ 0",
    "two-fraction-bars": "# ?/?/?",
    "condition-in-section-3": "[>1]0;[>2]0;[>3]0",
    "condition-in-section-3-with-text": "[>1]0;[>2]0;[>3]0;@",
    "condition-on-text-section": "0;0;0;[>1]@",
    "bare-letter-after-digits": "0 x",
    "bare-letter-before-digits": "x0",
    "condition-only-section": "[>1]",
    "color-only-section": "[Red]",
}


def build(directory: Path) -> list[str]:
    """One workbook per candidate; returns the file names in candidate order."""
    directory.mkdir(parents=True, exist_ok=True)
    for old in directory.glob("inv_*.xlsx"):
        old.unlink()
    names: list[str] = []
    for k, (label, fmt) in enumerate(CANDIDATES.items(), 1):
        BUILTIN_FORMATS_REVERSE.clear()
        wb = Workbook()
        ws = wb.active
        ws.title = "probe"
        for row, value in enumerate(VALUES, 1):
            cell = ws.cell(row=row, column=1, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.number_format = fmt
            ws.cell(row=row, column=2, value=label)
            ws.cell(row=row, column=3, value=repr(value))
            ws.cell(row=row, column=4, value=repr(fmt))
        name = f"inv_{k:02d}.xlsx"
        wb.save(directory / name)
        names.append(name)
    return names


def main() -> None:
    """Build, run, print."""
    directory = (WORKBOOK_DIR / "invalid").resolve()
    names = build(directory)
    proc = subprocess.run(
        ["osascript", str(SCRIPT), str(directory), str(len(VALUES))],
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode != 0:
        print(proc.stderr, file=sys.stderr)
        raise SystemExit(proc.returncode)
    by_name: dict[str, list[str]] = {}
    for line in proc.stdout.rstrip("\n").split("\n"):
        name, _, rest = line.partition("\t")
        by_name.setdefault(name, []).append(rest)
    for name, (label, fmt) in zip(names, CANDIDATES.items(), strict=True):
        lines = by_name.get(name, ["<no output>"])
        if lines[0].startswith("ERROR"):
            print(f"{label:36s} {fmt!r:26s} refused  ({lines[0][:40]})")
            continue
        shown = [ln.split("\t")[3] if ln.count("\t") >= 4 else ln for ln in lines]
        kept = lines[0].split("\t")[4] if lines[0].count("\t") >= 4 else "?"
        print(f"{label:36s} {fmt!r:26s} rendered {shown!r}  kept={kept!r}")


if __name__ == "__main__":
    main()
